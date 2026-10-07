"""取出測試數據：Excel（EXP-03、EXP-04）。"""

import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from flatness import export
from flatness.results import ResultSink, to_record
from flatness.store import MemoryStore
from test_results import T, result


def rec(serial, judgment="FAIL", at=T):
    r = to_record(result(serial, judgment), "ST01")
    r["head"]["measured_at"] = at
    return r


class WriteXlsxTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / "平整檢查_2026-10-07.xlsx"

    def test_columns_and_marks(self):  # EXP-03
        later = datetime(2026, 10, 7, 15, 0, 0)
        n = export.write_xlsx(self.path, [rec("20261007150000000", at=later), rec("20261007143005123")], T.date())
        self.assertEqual(n, 2)
        wb = load_workbook(self.path)
        ws = wb["量測紀錄"]
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("編號", "量測時間", "判定", "重讀次數", "前排 左", "前排 右", "前排 中", "不合格點"))
        first = rows[1]
        self.assertEqual(first[:4], ("20261007143005123", T, "不合格", 1))  # 依量測時間排序
        self.assertIsInstance(first[0], str)                                # 編號為文字
        self.assertEqual(ws.cell(2, 1).number_format, "@")
        self.assertEqual(first[4:7], (-0.2971, 12.5012, "異常"))
        self.assertEqual(first[7], "前排 左（偏低）、前排 中（異常）")
        self.assertEqual(ws.cell(2, 5).fill.fgColor.rgb[-6:], "FFC7CE")    # 偏低：紅底
        self.assertEqual(ws.cell(2, 6).fill.fill_type, None)              # 合格：無底色
        self.assertEqual(ws.cell(2, 7).fill.fgColor.rgb[-6:], "FFEB9C")    # 異常：黃底
        std = list(wb["當日標準值"].iter_rows(values_only=True))
        self.assertEqual(std[1][:6], ("第 1 排", 1, "前排 左", 12.5, 12.45, 12.55))
        self.assertEqual(std[3][3], "未設定")                              # 探頭 3 沒有標準
        self.assertEqual(list(Path(self.tmp.name).iterdir()), [self.path])  # 沒有留下暫存檔

    def test_no_records_no_file(self):  # EXP-04
        with self.assertRaises(ValueError):
            export.write_xlsx(self.path, [], T.date())
        self.assertFalse(self.path.exists())

    def test_unwritable_folder_raises(self):  # EXP-04：拔除隨身碟
        with self.assertRaises(OSError):
            export.write_xlsx(Path(self.tmp.name) / "gone" / "x.xlsx", [rec("20261007143005123")], T.date())


class SinkExportTest(unittest.TestCase):
    def test_db_plus_buffer(self):  # 資料庫 ＋ 尚未寫入的暫存，不重複
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        store = MemoryStore()
        r = rec("20261007143005123")
        store.write_inspection(r["head"], r["points"])
        sink = ResultSink(store, "ST01", Path(tmp.name) / "buffer")
        sink.dir.mkdir()
        for serial in ("20261007143005123", "20261007150000000"):  # 前者已在資料庫
            b = to_record(result(serial, "PASS"), "ST01")
            (sink.dir / f"{serial}.json").write_text(json.dumps(b, ensure_ascii=False))
        out = Path(tmp.name) / "out.xlsx"
        self.assertEqual(sink.export(T.date(), out), 2)
        serials = [row[0] for row in load_workbook(out)["量測紀錄"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(serials, ["20261007143005123", "20261007150000000"])


if __name__ == "__main__":
    unittest.main()
