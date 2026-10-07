"""取出測試數據：當日量測紀錄寫成 Excel（EXP-01～EXP-04、5.4）。

records 與 results.to_record 相同格式：{"head": {...}, "points": [...]}，measured_at 為 datetime。
工作表一「量測紀錄」一片一列：編號（文字）、量測時間、判定、重讀次數、各量測點測量值、不合格點；
偏高、偏低的儲存格紅底，設備異常的儲存格黃底並寫「異常」。工作表二「當日標準值」列出各點當日使用的
標準值與上下限（同一點當日改過標準時列出每一組與開始使用的時間）。
"""

from __future__ import annotations

import os
from datetime import date
from pathlib import Path

from .lan import row_label, row_no

JUDGMENT = {"PASS": "合格", "FAIL": "不合格", "ERROR": "設備異常"}
POINT_WORD = {"HIGH": "偏高", "LOW": "偏低", "ERROR": "異常"}
VALUE_FORMAT = "0.0000"


def write_xlsx(path, records: list[dict], day: date) -> int:
    """寫出 Excel 並回傳筆數；沒有紀錄時拋出 ValueError，不產生檔案（EXP-04）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if not records:
        raise ValueError(f"{day:%Y-%m-%d} 沒有紀錄")
    records = sorted(records, key=lambda r: r["head"]["measured_at"])
    names: dict[tuple, str] = {}
    for r in records:  # 量測點欄位：依排、探頭 ID；名稱取當日最後一次
        for p in r["points"]:
            names[(p["device_key"], p["probe_id"])] = f"{p['device_name']} {p['probe_description']}".strip()
    cols = sorted(names, key=lambda k: (row_no(k[0]) or 1 << 30, k[0], k[1]))

    red = PatternFill("solid", fgColor="FFC7CE")
    yellow = PatternFill("solid", fgColor="FFEB9C")
    bold = Font(bold=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "量測紀錄"
    ws.append(["編號", "量測時間", "判定", "重讀次數", *(names[k] for k in cols), "不合格點"])
    for r in records:
        h = r["head"]
        by = {(p["device_key"], p["probe_id"]): p for p in r["points"]}
        bad = [f"{names[k]}（{POINT_WORD[by[k]['judgment']]}）" for k in cols
               if k in by and by[k]["judgment"] in POINT_WORD]
        ws.append([h["serial"], h["measured_at"], JUDGMENT.get(h["judgment"], h["judgment"]), h["reread_count"],
                   *(None for _ in cols), "、".join(bad)])
        n = ws.max_row
        ws.cell(n, 1).number_format = "@"  # EXP-03：編號不被轉為數字或科學記號
        ws.cell(n, 2).number_format = "yyyy-mm-dd hh:mm:ss.000"
        if h["judgment"] != "PASS":
            ws.cell(n, 3).font = Font(bold=True, color="9C0006" if h["judgment"] == "FAIL" else "9C5700")
        for i, k in enumerate(cols, start=5):
            p = by.get(k)
            if p is None:
                continue
            c = ws.cell(n, i)
            if p["judgment"] == "ERROR" or p["measured_value"] is None:
                c.value, c.fill = "異常", yellow
                continue
            c.value, c.number_format = float(p["measured_value"]), VALUE_FORMAT
            if p["judgment"] in ("HIGH", "LOW"):
                c.fill = red
    for c in ws[1]:
        c.font, c.alignment = bold, Alignment(horizontal="center")
    widths = [20, 25, 10, 9, *(max(12, len(names[k]) * 2 + 2) for k in cols), 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"

    st = wb.create_sheet("當日標準值")
    st.append(["排", "探頭 ID", "量測點", "標準值", "下限", "上限", "開始使用時間"])
    for k in cols:
        seen = []
        for r in records:
            p = next((x for x in r["points"] if (x["device_key"], x["probe_id"]) == k), None)
            if p is None or p["standard_value"] is None:
                continue
            std = tuple(float(p[f]) for f in ("standard_value", "lower_limit", "upper_limit"))
            if not seen or seen[-1][0] != std:
                seen.append((std, r["head"]["measured_at"]))
        if not seen:
            st.append([row_label(k[0]), k[1], names[k], "未設定", None, None, None])
        for std, since in seen:
            st.append([row_label(k[0]), k[1], names[k], *std, since])
            n = st.max_row
            for i in (4, 5, 6):
                st.cell(n, i).number_format = "0.000"
            st.cell(n, 7).number_format = "yyyy-mm-dd hh:mm:ss"
    for c in st[1]:
        c.font = bold
    for i, w in enumerate([10, 8, 20, 10, 10, 10, 21], start=1):
        st.column_dimensions[get_column_letter(i)].width = w

    path = Path(path)
    tmp = path.with_name(f".{path.name}.tmp")
    try:  # 先寫暫存檔再改名：寫到一半失敗（例如拔除隨身碟）不留下不完整的檔案
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)
    return len(records)
